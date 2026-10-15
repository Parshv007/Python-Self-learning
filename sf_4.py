import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment


def create_batch_gradebook():
    # Step 1: Instantiate Workbook and set sheet title
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Semester_Results"

    # Step 2: Define and append Header Row
    headers = [
        "Roll No", 
        "Student Name", 
        "Lab Marks (out of 50)", 
        "Theory Marks (out of 50)", 
        "Total Marks (out of 100)", 
        "Result"
    ]
    ws.append(headers)

    # Step 3: Raw sample student records [Roll No, Name, Lab Marks, Theory Marks]
    students_data = [
        [101, "Aarav Sharma", 42, 38],
        [102, "Ananya Patel", 18, 15],
        [103, "Rohan Verma", 45, 47],
        [104, "Priya Nair", 25, 20],
        [105, "Kabir Mehta", 35, 30]
    ]

    # Step 4: Add student rows and perform calculations dynamically
    total_class_marks = 0

    for student in students_data:
        roll_no, name, lab, theory = student
        
        # Calculate Total Marks and Result
        total = lab + theory
        result = "Pass" if total >= 40 else "Fail"
        
        # Keep running tally for class average
        total_class_marks += total
        
        # Append completed row to worksheet
        ws.append([roll_no, name, lab, theory, total, result])

    # Calculate Class Average
    num_students = len(students_data)
    class_average = total_class_marks / num_students if num_students > 0 else 0

    # Append summary row
    ws.append([])  # Blank spacer row
    summary_row = ["Class Average", "", "", "", round(class_average, 2), ""]
    ws.append(summary_row)

    # Step 5: Apply Styling (Header bolding, gray fill, centered alignment)
    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid") # Light Gray
    header_font = Font(bold=True)
    summary_font = Font(bold=True)

    # Style Header Row (Row 1)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Style Summary Row (Last Row)
    last_row_idx = ws.max_row
    for cell in ws[last_row_idx]:
        cell.font = summary_font

    # Save the workbook to disk
    output_filename = "AIML_Batch_Grades.xlsx"
    wb.save(output_filename)
    print(f"Success! Gradebook generated and saved to '{output_filename}'.")


if __name__ == "__main__":
    create_batch_gradebook()